"""Editable Word manuscript mirroring the print interior page-for-page (same trim + bleed size,
same page order, illustrations on their own pages)."""
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, sys.argv[1])
import content as C  # noqa: E402
ILL, OUT = sys.argv[2], sys.argv[3]

NAVY = RGBColor(0x1B, 0x2A, 0x4A); GOLD = RGBColor(0x8A, 0x6A, 0x1F); MAROON = RGBColor(0x7A, 0x24, 0x18)
INK = RGBColor(0x2A, 0x25, 0x22)
BODY, HEAD = 'Literata', 'Cinzel'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.625), Inches(11.25)
# mirror margins (inside/outside) like the print file; values measured from the page edge
sec.top_margin, sec.bottom_margin = Inches(0.825), Inches(0.875)
sec.left_margin, sec.right_margin = Inches(0.875), Inches(0.875)
settings = doc.settings.element
mm = OxmlElement('w:mirrorMargins'); settings.append(mm)

st = doc.styles['Normal']
st.font.name = BODY; st.font.size = Pt(11); st.font.color.rgb = INK
st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY)
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.15


def para(text='', size=None, bold=False, italic=False, color=None, font=None, align=None, after=None, before=None, caps=False):
    p = doc.add_paragraph()
    if BREAK[0]:
        p.paragraph_format.page_break_before = True; BREAK[0] = False
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        r.font.all_caps = caps
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        if font:
            r.font.name = font; r._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    return p


def shade(p, hexfill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexfill)
    pPr.append(shd)


BREAK = [False]


def page_break():
    BREAK[0] = True


def art_page(idx):
    p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    p.add_run().add_picture(f'{ILL}/ill-{idx:02d}.jpg', height=Inches(9.5))


C_ = WD_ALIGN_PARAGRAPH.CENTER
# 1 Title
para('', after=90)
para('JOESIAH', 44, True, color=NAVY, font=HEAD, align=C_, after=0)
para('&', 30, True, color=GOLD, font=HEAD, align=C_, after=0)
para('MALANI', 44, True, color=NAVY, font=HEAD, align=C_, after=18)
para('NOAH’S BLOODLINE', 26, True, color=MAROON, font=HEAD, align=C_, after=4)
para('BOOK 1', 12, True, color=GOLD, font=HEAD, align=C_, after=18)
para(C.TAGLINE, 14, italic=True, color=NAVY, align=C_, after=150)
para('Written by', 12, align=C_, after=2)
para(C.AUTHOR.upper(), 18, True, color=NAVY, font=HEAD, align=C_)
page_break()
# 2 Copyright
para('', after=300)
for t, kw in [('Joesiah & Malani: Noah’s Bloodline', dict(size=11, bold=True, color=NAVY, font=HEAD)),
              ('Joesiah & Malani, Book 1', dict(size=9.5)),
              ('Text copyright © 2026 %s\nAll rights reserved.' % C.AUTHOR, dict(size=9.5)),
              ('No part of this book may be reproduced, stored in a retrieval system, or transmitted in any form or by any means—electronic, mechanical, photocopying, recording, or otherwise—without the prior written permission of the author, except for brief quotations used in reviews.', dict(size=9.5)),
              ('Unless otherwise noted, Scripture quotations in the story text are taken from the Holy Bible, New International Version®, NIV®. Copyright © 1973, 1978, 1984, 2011 by Biblica, Inc.® Used by permission. All rights reserved worldwide. Scripture quotations within the illustrations are drawn from or adapted from several English translations.', dict(size=9.5)),
              ('First edition, 2026', dict(size=9.5))]:
    para(t, align=C_, after=8, **kw)
page_break()
# 3 Contents
PAGE_OF_CH = {n: 2 * n + 2 for n in range(1, 12)}; PAGE_OF_CH.update({12: 26, 13: 28, 14: 30})
para('CONTENTS', 24, True, color=NAVY, font=HEAD, align=C_, after=20)
for n, t, _ in C.CHAPTERS:
    p = para(after=9)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.85), alignment=2, leader=1)  # right, dots
    r = p.add_run('%d   ' % n); r.font.color.rgb = GOLD; r.bold = True
    p.add_run(t + '\t' + str(PAGE_OF_CH[n]))
page_break()


def chapter(n, title, blocks):
    para('CHAPTER %d' % n, 11, True, color=GOLD, font=HEAD, align=C_, after=2)
    para(title, 19, True, color=NAVY, font=HEAD, align=C_, after=16)
    for b in blocks:
        k = b[0]
        if k == 'p':
            para(b[1])
        elif k == 'quote':
            p = para(b[1], italic=True, color=NAVY, after=0); p.paragraph_format.left_indent = Inches(0.35); p.paragraph_format.right_indent = Inches(0.35)
            p = para('—' + b[2], 8.5, True, color=GOLD, font=HEAD, align=WD_ALIGN_PARAGRAPH.RIGHT); p.paragraph_format.right_indent = Inches(0.35)
        elif k in ('names', 'emph'):
            txt = '  •  '.join(x.upper() for x in b[1]) if k == 'names' else b[1].upper()
            para(txt, 13.5, True, color=NAVY, font=HEAD, align=C_)
        elif k == 'motto':
            p = para(C.MOTTO, 11, italic=True, color=NAVY, align=C_, before=4, after=8); shade(p, 'F3E7C9')
        elif k == 'thought':
            if b[1]:
                para(b[1], 8.5, True, color=GOLD, font=HEAD, after=0)
            p = para(b[2], bold=True, italic=True, color=NAVY); shade(p, 'F7EED8')
        elif k == 'pledge':
            para(b[1], 13, True, color=MAROON, font=HEAD, align=C_)
        elif k == 'turn':
            p = para(align=WD_ALIGN_PARAGRAPH.RIGHT)
            r = p.add_run('TURN THE PAGE →  '); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = GOLD
            r = p.add_run(b[1]); r.italic = True; r.font.color.rgb = NAVY
        elif k == 'read':
            para(b[1], 8.5, True, color=GOLD, font=HEAD, align=C_, before=14, after=0)
            para(b[2], 10.5, align=C_)
        elif k == 'end':
            for t, kw in [('THE JOURNEY CONTINUES', dict(size=11, bold=True, color=GOLD, font=HEAD)),
                          ('Joesiah & Malani • Book 2', dict(size=12, bold=True, color=NAVY, font=HEAD)),
                          (C.NEXT_BOOK, dict(italic=True, color=NAVY)),
                          (C.MOTTO, dict(size=10.5, italic=True, color=NAVY))]:
                p = para(t, align=C_, after=3, **kw); shade(p, 'F3E7C9')


chap = {n: (t, b) for n, t, b in C.CHAPTERS}
for n in range(1, 15):
    chapter(n, *chap[n]); page_break()
    if n == 12:
        para('', after=150)
        para(C.SPOTLIGHT['label'], 13, True, color=GOLD, font=HEAD, align=C_)
        para(C.SPOTLIGHT['intro'], 12, italic=True, align=C_, after=14)
        para(C.SPOTLIGHT['text'], 20, italic=True, color=NAVY, align=C_, after=14)
        para('—' + C.SPOTLIGHT['ref'].upper(), 13, True, color=GOLD, font=HEAD, align=C_)
        page_break()
        continue
    art_page(C.CHAPTER_ART[n])
    if n < 14:
        page_break()

doc.core_properties.title = '%s: %s' % (C.TITLE, C.SUBTITLE)
doc.core_properties.author = C.AUTHOR
doc.core_properties.subject = '%s, Book %d' % (C.SERIES, C.BOOK_NO)
doc.save(OUT)
print('saved', OUT)
